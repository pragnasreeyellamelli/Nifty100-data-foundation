from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from .engine import (
    load_screener_data,
    run_preset_screener,
    PRESET_SCREENERS,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DIR = PROJECT_ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "screener_output.xlsx"


GREEN_FILL = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE",
)

RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE",
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)


def _sheet_name(name):
    invalid = '[]:*?/\\'

    for char in invalid:
        name = name.replace(char, "")

    return name[:31]


def _preset_rules(preset_name):
    preset = PRESET_SCREENERS[preset_name]

    if preset_name == "Quality Compounder":
        return [
            ("return_on_equity_pct", preset["roe_min"], "min"),
            ("debt_to_equity", preset["de_max"], "max"),
            ("free_cash_flow_cr", preset["fcf_min"], "min"),
            ("revenue_cagr_5yr", preset["revenue_cagr_5yr_min"], "min"),
        ]

    if preset_name == "Value Pick":
        return [
            ("pe_ratio", preset["pe_max"], "max"),
            ("pb_ratio", preset["pb_max"], "max"),
            ("debt_to_equity", preset["de_max"], "max"),
            ("dividend_yield_pct", preset["dividend_yield_min"], "min"),
        ]

    if preset_name == "Growth Accelerator":
        return [
            ("pat_cagr_5yr", preset["pat_cagr_5yr_min"], "min"),
            ("revenue_cagr_5yr", preset["revenue_cagr_5yr_min"], "min"),
            ("debt_to_equity", preset["de_max"], "max"),
        ]

    if preset_name == "Dividend Champion":
        return [
            ("dividend_yield_pct", preset["dividend_yield_min"], "min"),
            ("dividend_payout", preset["dividend_payout_max"], "max"),
            ("free_cash_flow_cr", preset["fcf_min"], "min"),
        ]

    if preset_name == "Debt-Free Blue Chip":
        return [
            ("debt_to_equity", preset["de_max"], "max"),
            ("return_on_equity_pct", preset["roe_min"], "min"),
            ("sales", preset["sales_min"], "min"),
        ]

    if preset_name == "Turnaround Watch":
        return [
            ("revenue_cagr_3yr", preset["revenue_cagr_3yr_min"], "min"),
            ("free_cash_flow_cr", preset["fcf_min"], "min"),
            ("de_declining", True, "boolean"),
        ]

    return []


def _apply_cell_colors(sheet, dataframe, rules):

    column_numbers = {
        column: index + 1
        for index, column in enumerate(dataframe.columns)
    }

    for metric, threshold, comparison in rules:

        if metric not in column_numbers:
            continue

        column_number = column_numbers[metric]

        for row_number, (_, row) in enumerate(
            dataframe.iterrows(),
            start=2,
        ):

            value = row[metric]

            if pd.isna(value):
                continue

            if comparison == "min":
                try:
                    passed = float(value) >= float(threshold)
                except (TypeError, ValueError):
                    continue

            elif comparison == "max":
                try:
                    passed = float(value) <= float(threshold)
                except (TypeError, ValueError):
                    continue

            elif comparison == "boolean":
                passed = bool(value)

            else:
                continue

            sheet.cell(
                row=row_number,
                column=column_number,
            ).fill = (
                GREEN_FILL
                if passed
                else RED_FILL
            )


def export_screener_output(output_file=OUTPUT_FILE):
    """
    Generate output/screener_output.xlsx
    with six preset worksheets.
    """

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    df = load_screener_data()

    workbook = Workbook()

    workbook.remove(
        workbook.active
    )

    for preset_name in PRESET_SCREENERS:

        result = run_preset_screener(
            df,
            preset_name,
        )

        if "composite_quality_score" in result.columns:
            result = result.sort_values(
                "composite_quality_score",
                ascending=False,
                kind="stable",
            )

        preferred_columns = [
            "company_id",
            "ticker",
            "company_name",
            "sector",
            "return_on_equity_pct",
            "return_on_capital_employed_pct",
            "net_profit_margin_pct",
            "debt_to_equity",
            "interest_coverage",
            "free_cash_flow_cr",
            "cash_from_operations_cr",
            "revenue_cagr_5yr",
            "pat_cagr_5yr",
            "eps_cagr_5yr",
            "operating_profit_margin_pct",
            "pe_ratio",
            "pb_ratio",
            "dividend_yield_pct",
            "dividend_payout",
            "sales",
            "net_profit",
            "asset_turnover",
            "composite_quality_score",
        ]

        columns = [
            column
            for column in preferred_columns
            if column in result.columns
        ]

        export_df = result[columns].copy()

        sheet = workbook.create_sheet(
            title=_sheet_name(preset_name)
        )

        # Header
        for column_number, column_name in enumerate(
            export_df.columns,
            start=1,
        ):

            cell = sheet.cell(
                row=1,
                column=column_number,
                value=column_name,
            )

            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # Data
        for row_number, (_, row) in enumerate(
            export_df.iterrows(),
            start=2,
        ):

            for column_number, column_name in enumerate(
                export_df.columns,
                start=1,
            ):

                value = row[column_name]

                if pd.isna(value):
                    value = None

                sheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )

        # Threshold colours
        _apply_cell_colors(
            sheet,
            export_df,
            _preset_rules(preset_name),
        )

        # Formatting
        sheet.freeze_panes = "A2"

        if sheet.max_row >= 1:
            sheet.auto_filter.ref = sheet.dimensions

        sheet.row_dimensions[1].height = 25

        for column_cells in sheet.columns:

            column_letter = get_column_letter(
                column_cells[0].column
            )

            maximum = 0

            for cell in column_cells:

                if cell.value is not None:
                    maximum = max(
                        maximum,
                        len(str(cell.value)),
                    )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max(maximum + 2, 12),
                28,
            )

    workbook.save(output_file)

    return output_file


if __name__ == "__main__":
    print(
        export_screener_output()
    )