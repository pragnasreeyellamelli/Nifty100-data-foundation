from pathlib import Path
import sqlite3

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]

DB_PATH = PROJECT_ROOT / "nifty100.db"
OUTPUT_DIR = PROJECT_ROOT / "output"
OUTPUT_FILE = OUTPUT_DIR / "peer_comparison.xlsx"


# ============================================================
# STYLES
# ============================================================

GREEN_FILL = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE",
)

YELLOW_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFEB9C",
)

RED_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE",
)

BENCHMARK_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFD966",
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
)

MEDIAN_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

MEDIAN_FONT = Font(
    bold=True,
)


# ============================================================
# HELPERS
# ============================================================

def _safe_sheet_name(name):
    name = str(name)

    for char in '[]:*?/\\':
        name = name.replace(char, "")

    return name[:31]


def _numeric(series):
    return pd.to_numeric(
        series,
        errors="coerce",
    )


# ============================================================
# LOAD DATA
# ============================================================

def load_percentiles():

    with sqlite3.connect(DB_PATH) as conn:

        df = pd.read_sql_query(
            """
            SELECT
                company_id,
                peer_group_name,
                metric,
                value,
                percentile_rank,
                year
            FROM peer_percentiles
            """,
            conn,
        )

    df["company_id"] = (
        df["company_id"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )

    df["peer_group_name"] = (
        df["peer_group_name"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    return df


def load_company_names():

    with sqlite3.connect(DB_PATH) as conn:

        companies = pd.read_sql_query(
            """
            SELECT
                company_id,
                company_name
            FROM companies
            """,
            conn,
        )

    companies["company_id"] = (
        companies["company_id"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )

    return companies


def load_benchmarks():

    path = (
        PROJECT_ROOT
        / "data"
        / "raw"
        / "peer_groups.xlsx"
    )

    peer = pd.read_excel(
        path
    )

    peer["company_id"] = (
        peer["company_id"]
        .fillna("")
        .astype(str)
        .str.strip()
        .str.upper()
    )

    peer["peer_group_name"] = (
        peer["peer_group_name"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    if "is_benchmark" not in peer.columns:
        peer["is_benchmark"] = False

    peer["is_benchmark"] = (
        peer["is_benchmark"]
        .fillna(False)
        .astype(bool)
    )

    return peer[
        [
            "company_id",
            "peer_group_name",
            "is_benchmark",
        ]
    ]


# ============================================================
# BUILD ONE PEER GROUP
# ============================================================

def build_group_table(
    group_name,
    percentiles,
    company_names,
    benchmarks,
):
    """
    Build wide-format table:
      company_id
      company_name
      10 metric values
      10 percentile columns
    """

    group = percentiles[
        percentiles["peer_group_name"]
        == group_name
    ].copy()

    if group.empty:
        return pd.DataFrame()

    # --------------------------------------------------------
    # Add company names
    # --------------------------------------------------------

    group = group.merge(
        company_names,
        on="company_id",
        how="left",
    )

    group["company_name"] = (
        group["company_name"]
        .fillna("")
        .astype(str)
        .str.strip()
    )

    # --------------------------------------------------------
    # Metric values
    # --------------------------------------------------------

    values = group.pivot_table(
        index=[
            "company_id",
            "company_name",
        ],
        columns="metric",
        values="value",
        aggfunc="first",
    ).reset_index()

    # --------------------------------------------------------
    # Percentiles
    # --------------------------------------------------------

    ranks = group.pivot_table(
        index=[
            "company_id",
            "company_name",
        ],
        columns="metric",
        values="percentile_rank",
        aggfunc="first",
    ).reset_index()

    ranks = ranks.rename(
        columns={
            column: f"{column} Percentile"
            for column in ranks.columns
            if column
            not in [
                "company_id",
                "company_name",
            ]
        }
    )

    result = values.merge(
        ranks,
        on=[
            "company_id",
            "company_name",
        ],
        how="outer",
    )

    # --------------------------------------------------------
    # Benchmark
    # --------------------------------------------------------

    group_benchmarks = benchmarks[
        benchmarks["peer_group_name"]
        == group_name
    ][
        [
            "company_id",
            "is_benchmark",
        ]
    ]

    result = result.merge(
        group_benchmarks,
        on="company_id",
        how="left",
    )

    result["is_benchmark"] = (
        result["is_benchmark"]
        .fillna(False)
        .astype(bool)
    )

    return result


# ============================================================
# PERCENTILE COLOURS
# ============================================================

def colour_percentiles(sheet):

    for column_number in range(
        1,
        sheet.max_column + 1,
    ):

        header = sheet.cell(
            row=1,
            column=column_number,
        ).value

        if (
            header is None
            or "Percentile"
            not in str(header)
        ):
            continue

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):

            value = sheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if not isinstance(
                value,
                (int, float),
            ):
                continue

            if value >= 75:
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).fill = GREEN_FILL

            elif value <= 25:
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).fill = RED_FILL

            else:
                sheet.cell(
                    row=row_number,
                    column=column_number,
                ).fill = YELLOW_FILL


# ============================================================
# MEDIAN ROW
# ============================================================

def add_median_row(
    sheet,
    dataframe,
):
    """
    Add peer-group median for every numeric column.
    """

    row_number = sheet.max_row + 1

    sheet.cell(
        row=row_number,
        column=1,
        value="Peer Group Median",
    )

    sheet.merge_cells(
        start_row=row_number,
        start_column=1,
        end_row=row_number,
        end_column=2,
    )

    for column_number in range(
        1,
        sheet.max_column + 1,
    ):

        sheet.cell(
            row=row_number,
            column=column_number,
        ).fill = MEDIAN_FILL

        sheet.cell(
            row=row_number,
            column=column_number,
        ).font = MEDIAN_FONT

    for column_number in range(
        3,
        sheet.max_column + 1,
    ):

        header = sheet.cell(
            row=1,
            column=column_number,
        ).value

        if header is None:
            continue

        if header in dataframe.columns:

            values = _numeric(
                dataframe[header]
            )

        elif (
            str(header)
            .endswith(" Percentile")
        ):

            values = _numeric(
                dataframe[header]
            )

        else:
            continue

        values = values.dropna()

        if not values.empty:

            sheet.cell(
                row=row_number,
                column=column_number,
                value=float(
                    values.median()
                ),
            )


# ============================================================
# EXPORT
# ============================================================

def export_peer_comparison(
    output_file=OUTPUT_FILE,
):
    """
    Generate exactly 11 peer-group worksheets.
    """

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    percentiles = load_percentiles()

    company_names = load_company_names()

    benchmarks = load_benchmarks()

    real_groups = sorted(
        group
        for group in
        percentiles[
            "peer_group_name"
        ].dropna().unique()
        if group
        and group
        != "No peer group assigned"
    )

    print(
        "PEER GROUPS FOUND:",
        real_groups,
    )

    workbook = Workbook()

    workbook.remove(
        workbook.active
    )

    created_groups = []

    for group_name in real_groups:

        table = build_group_table(
            group_name,
            percentiles,
            company_names,
            benchmarks,
        )

        if table.empty:
            continue

        created_groups.append(
            group_name
        )

        # ----------------------------------------------------
        # Remove helper column
        # ----------------------------------------------------

        is_benchmark = table[
            "is_benchmark"
        ].copy()

        export_df = table.drop(
            columns=[
                "is_benchmark"
            ],
            errors="ignore",
        ).copy()

        # ----------------------------------------------------
        # Worksheet
        # ----------------------------------------------------

        sheet = workbook.create_sheet(
            title=_safe_sheet_name(
                group_name
            )
        )

        # ----------------------------------------------------
        # Header
        # ----------------------------------------------------

        for column_number, column_name in enumerate(
            export_df.columns,
            start=1,
        ):

            cell = sheet.cell(
                row=1,
                column=column_number,
                value=column_name,
            )

            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        # ----------------------------------------------------
        # Data
        # ----------------------------------------------------

        for row_number, (_, row) in enumerate(
            export_df.iterrows(),
            start=2,
        ):

            for column_number, column_name in enumerate(
                export_df.columns,
                start=1,
            ):

                value = row[column_name]

                if pd.isna(value):
                    value = None

                sheet.cell(
                    row=row_number,
                    column=column_number,
                    value=value,
                )

            # Benchmark highlight.
            if bool(
                is_benchmark.iloc[
                    row_number - 2
                ]
            ):

                for column_number in range(
                    1,
                    sheet.max_column + 1,
                ):

                    sheet.cell(
                        row=row_number,
                        column=column_number,
                    ).fill = (
                        BENCHMARK_FILL
                    )

        # ----------------------------------------------------
        # Percentile colours
        # ----------------------------------------------------

        colour_percentiles(
            sheet
        )

        # ----------------------------------------------------
        # Median row
        # ----------------------------------------------------

        add_median_row(
            sheet,
            export_df,
        )

        # ----------------------------------------------------
        # Formatting
        # ----------------------------------------------------

        sheet.freeze_panes = "A2"

        sheet.row_dimensions[
            1
        ].height = 25

        sheet.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(sheet.max_column)}"
            f"{sheet.max_row - 1}"
        )

        for column_cells in sheet.columns:

            letter = get_column_letter(
                column_cells[0].column
            )

            max_length = 0

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value)),
                    )

            sheet.column_dimensions[
                letter
            ].width = min(
                max(
                    max_length + 2,
                    12,
                ),
                28,
            )

    # --------------------------------------------------------
    # Verify exactly 11 sheets
    # --------------------------------------------------------

    print(
        "SHEETS CREATED:",
        created_groups,
    )

    if len(workbook.sheetnames) != 11:

        raise ValueError(
            "Expected exactly 11 peer-group sheets, "
            f"but created {len(workbook.sheetnames)}."
        )

    workbook.save(
        output_file
    )

    return output_file


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    output = export_peer_comparison()

    print(
        "PEER COMPARISON CREATED:"
    )

    print(output)